"""
Training pipeline — Stage 2 & 3 from the product spec (Section 5):
parse an uploaded presentation file into structured per-slide data,
then label each slide's category via AI. This is what turns a raw
uploaded file into actual usable training examples, instead of just
sitting in storage.

Stage 4 (accumulate ~500-1000 examples per category) and Stage 5
(train a small classifier on the accumulated examples) come after
this — they need real volume, which this pipeline is what generates.
"""

import os
from io import BytesIO

MAX_SLIDES_PER_UPLOAD = 50  # caps AI-labeling cost per upload


def parse_pptx(file_bytes):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    slides_data = []
    prs = Presentation(BytesIO(file_bytes))
    for slide in prs.slides:
        texts = []
        bullet_count = 0
        has_image = False
        for shape in slide.shapes:
            if shape.shape_type == MSO_SHAPE_TYPE.PICTURE:
                has_image = True
            if shape.has_text_frame:
                for para in shape.text_frame.paragraphs:
                    text = "".join(run.text for run in para.runs)
                    if text.strip():
                        texts.append(text.strip())
                        bullet_count += 1
        slides_data.append({
            "raw_text": "\n".join(texts)[:2000],  # cap stored text size
            "bullet_count": bullet_count,
            "has_image": has_image,
        })
    return slides_data


def parse_pdf(file_bytes):
    """
    Renders each page to an image and reads it via vision-based
    extraction (Gemini, falling back to Tesseract OCR), instead of
    pulling the PDF's embedded text layer directly.

    Why: many Bangla PDFs (including ones built by rendering text with
    a custom font and embedding it) don't carry a proper Unicode
    mapping for their text layer — pulling text directly yields
    garbled control-character soup even though the page displays
    correctly. Reading the rendered image sidesteps that entirely.
    """
    import pymupdf as fitz  # PyMuPDF (using new import name to avoid deprecation warning)
    from ocr import extract_text_from_image_bytes
    from vision_ocr import extract_content_from_image
    import tempfile

    slides_data = []
    doc = fitz.open(stream=file_bytes, filetype="pdf")
    for page in doc:
        # 100 dpi (not 150) — cuts memory per page substantially; plenty
        # of resolution for OCR/vision on typical exam/notes text, and
        # this service runs under a tight 512MB RAM limit.
        pix = page.get_pixmap(dpi=100)
        img_bytes = pix.tobytes("png")

        text = extract_content_from_image(img_bytes, mime_type="image/png")
        if text is None:
            text = extract_text_from_image_bytes(img_bytes)

        text = (text or "").strip()
        bullet_count = len([l for l in text.split("\n") if l.strip()])
        # a page that's mostly a diagram/photo with little extracted
        # text is a reasonable proxy for "has an image worth flagging"
        has_image = bullet_count <= 2
        slides_data.append({
            "raw_text": text[:2000],
            "bullet_count": bullet_count,
            "has_image": has_image,
        })
        pix = None  # release native pixmap memory promptly, don't wait for GC
    doc.close()
    return slides_data


def parse_file(filename, file_bytes):
    """Returns a list of slide-dicts, or None if the format isn't supported yet."""
    ext = filename.lower().rsplit(".", 1)[-1] if "." in filename else ""
    if ext == "pptx":
        return parse_pptx(file_bytes)
    if ext == "pdf":
        return parse_pdf(file_bytes)
    return None


def label_slide(raw_text, bullet_count, has_image):
    """
    Classifies a slide into a category. Uses Gemini if configured;
    falls back to a simple heuristic otherwise (or if the call fails)
    so the pipeline never breaks just because labeling is unavailable —
    it just produces slightly less precise labels.
    """
    if os.environ.get("GEMINI_API_KEY"):
        try:
            from google import genai
            import gemini_client

            client = genai.Client(api_key=os.environ["GEMINI_API_KEY"])
            prompt = (
                "Classify this slide's content into exactly ONE of these categories: "
                "title, bullet-list, image-heavy, comparison, notes-formula, other. "
                "Return ONLY the category word, nothing else.\n\n"
                f"Text: {raw_text[:500]}\nBullet count: {bullet_count}\nHas image: {has_image}"
            )
            resp = gemini_client.generate_content(client, contents=prompt)
            label = resp.text.strip().lower().split()[0] if resp.text.strip() else ""
            valid = {"title", "bullet-list", "image-heavy", "comparison", "notes-formula", "other"}
            if label in valid:
                return label
        except Exception:
            pass

    # heuristic fallback (also what runs if GEMINI_API_KEY isn't set)
    if has_image and bullet_count <= 1:
        return "image-heavy"
    if bullet_count <= 1:
        return "title"
    if bullet_count >= 4:
        return "bullet-list"
    return "other"


def process_upload(filename, file_bytes):
    """
    Full pipeline for one uploaded file: parse -> label each slide
    (capped at MAX_SLIDES_PER_UPLOAD) -> return structured results.
    Returns (status, parsed_slides) where status is 'processed' or
    'unsupported_format' or 'failed'.
    """
    try:
        slides = parse_file(filename, file_bytes)
    except Exception:
        return "failed", []

    if slides is None:
        return "unsupported_format", []

    results = []
    for slide in slides[:MAX_SLIDES_PER_UPLOAD]:
        category = label_slide(slide["raw_text"], slide["bullet_count"], slide["has_image"])
        results.append({**slide, "category": category})

    return "processed", results
